#!/usr/bin/env python3
"""Build catalog.js and item photos from the Estancia La Taba spreadsheet."""

import json
import re
import unicodedata
import zipfile
from collections import defaultdict
from io import BytesIO
from pathlib import Path

import openpyxl
from PIL import Image, ImageOps

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "data" / "inventario.xlsx"
HERO_SRC = Path(
    "/Users/RP333/.cursor/projects/Users-RP333-Inventsrio-La-Taba/assets/"
    "Captura_de_Pantalla_2026-09-25_a_la_s__11.04.03-9d620374-a4df-4d01-9fac-2b8aaa5cb59e.jpg"
)
ITEMS_DIR = ROOT / "images" / "items"
HERO_DIR = ROOT / "images" / "hero"

ROOM_LABELS = {
    "Resibidor": "Recibidor",
    "Salon de juegos": "Salón de juegos",
    "Sector Bar": "Sector bar",
    "Sector bar": "Sector bar",
    "Fogon": "Fogón",
    "Fogon ": "Fogón",
    "Fogón": "Fogón",
}

CATEGORY_LABELS = {
    "Electrodimesticos": "Electrodomésticos",
    "Iluminacion": "Iluminación",
    "vajilla": "Vajilla",
    "Cuadro animales": "Cuadros de animales",
    "Cuadros caballos": "Cuadros de caballos",
    "Fogon": "Fogón",
    "Fogon ": "Fogón",
}


def slug(value):
    text = unicodedata.normalize("NFKD", value)
    text = text.encode("ascii", "ignore").decode()
    text = re.sub(r"[^a-z0-9]+", "-", text.lower()).strip("-")
    return text or "ambiente"


def label(value, table):
    if value is None:
        return None
    text = str(value).strip()
    return table.get(text, text)


def load_images_by_row():
    with zipfile.ZipFile(XLSX) as workbook:
        rels = workbook.read("xl/drawings/_rels/drawing1.xml.rels").decode()
        rid_to_img = dict(
            re.findall(r'Id="(rId\d+)"[^>]*Target="\.\./media/([^"]+)"', rels)
        )
        drawing = workbook.read("xl/drawings/drawing1.xml").decode()
        anchors = re.findall(
            r'<xdr:row>(\d+)</xdr:row>.*?<a:blip[^>]*r:embed="(rId\d+)"',
            drawing,
            re.S,
        )
        blobs = {}
        by_row = {}
        for row, rid in anchors:
            name = rid_to_img[rid]
            if name not in blobs:
                blobs[name] = workbook.read(f"xl/media/{name}")
            by_row[int(row) + 1] = blobs[name]
        return by_row


def catalog_rows(images):
    workbook = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = workbook["CATALOGO FERIA"]
    section = None
    rows = []
    for index, row in enumerate(
        sheet.iter_rows(min_row=6, max_row=220, max_col=5, values_only=True), 6
    ):
        image_cell, name, _description, _nuevo, _usado = row
        if name:
            rows.append(
                {
                    "room": section,
                    "image": images.get(index),
                    "needs_photo": str(image_cell or "").strip().upper() == "AGREGAR FOTO",
                }
            )
        elif image_cell and not name:
            header = str(image_cell).strip()
            if header.upper() != "AGREGAR FOTO":
                section = header
    return rows


def inventory_rows():
    workbook = openpyxl.load_workbook(XLSX, data_only=True)
    sheet = workbook["INVENTARIO"]
    items = []
    for row in sheet.iter_rows(min_row=2, max_row=220, max_col=12, values_only=True):
        if row[0] and row[3]:
            items.append(row)
    return items


def save_photo(blob, destination):
    image = ImageOps.exif_transpose(Image.open(BytesIO(blob)))
    if image.mode != "RGB":
        image = image.convert("RGB")
    image.thumbnail((1600, 1600), Image.Resampling.LANCZOS)
    image.save(destination, "JPEG", quality=84, optimize=True)


def cover_pins(image):
    # The screenshot includes small white sheet pins sitting in the sky.
    pixels = image.load()
    width, _height = image.size
    seeds = []
    for y in range(10, 42):
        for x in range(width):
            red, green, blue = pixels[x, y]
            if red > 245 and green > 245 and blue > 240:
                seeds.append((x, y))
    mask = set(seeds)
    for x, y in seeds:
        for dy in (-1, 0, 1):
            for dx in (-1, 0, 1):
                nx, ny = x + dx, y + dy
                if not (0 <= nx < width and 10 <= ny < 42):
                    continue
                red, green, blue = pixels[nx, ny]
                if red > 220 and green > 220 and blue > 210:
                    mask.add((nx, ny))
    for x, y in mask:
        sample = None
        for sample_y in (5, 6, 4, 7, 8):
            if (x, sample_y) in mask:
                continue
            red, green, blue = pixels[x, sample_y]
            if blue > 190 and red > 170:
                sample = (red, green, blue)
                break
        if sample is None:
            sample = pixels[max(0, x - 24), 6]
        pixels[x, y] = sample
    return image


def export_hero():
    HERO_DIR.mkdir(parents=True, exist_ok=True)
    image = cover_pins(Image.open(HERO_SRC).convert("RGB"))
    panels = [
        (0, 3, 339, image.height),
        (340, 3, 680, image.height),
        (681, 3, 1021, image.height),
    ]
    names = ["izquierda.jpg", "centro.jpg", "derecha.jpg"]
    for box, name in zip(panels, names):
        panel = image.crop(box)
        panel.save(HERO_DIR / name, "JPEG", quality=90, optimize=True)
        print(name, panel.size)


def main():
    images = load_images_by_row()
    catalog = catalog_rows(images)
    inventory = inventory_rows()
    if len(catalog) != len(inventory):
        raise SystemExit(f"Row mismatch: catalog {len(catalog)} inventory {len(inventory)}")

    ITEMS_DIR.mkdir(parents=True, exist_ok=True)
    for old in ITEMS_DIR.glob("*"):
        old.unlink()

    items = []
    room_order = []
    missing = 0
    for photo, row in zip(catalog, inventory):
        item_id, ambiente, categoria, nombre, descripcion, cantidad, medidas, estado, precio, venta, publicar, *_rest = row
        room = label(photo["room"] or ambiente, ROOM_LABELS)
        if room not in room_order:
            room_order.append(room)
        image_path = None
        images = []
        manual_many = sorted((ROOT / "images" / "manual").glob(f"{item_id}-*.jpg"))
        manual_one = ROOT / "images" / "manual" / f"{item_id}.jpg"
        if manual_many:
            for index, manual in enumerate(manual_many, start=1):
                relative = f"images/items/{item_id}-{index}.jpg"
                save_photo(manual.read_bytes(), ROOT / relative)
                images.append(relative)
            image_path = images[0]
        elif manual_one.exists():
            image_path = f"images/items/{item_id}.jpg"
            save_photo(manual_one.read_bytes(), ROOT / image_path)
        elif photo["image"] and not photo["needs_photo"]:
            image_path = f"images/items/{item_id}.jpg"
            save_photo(photo["image"], ROOT / image_path)
        else:
            missing += 1
        quantity = int(cantidad) if isinstance(cantidad, (int, float)) else 1
        price = int(precio) if isinstance(precio, (int, float)) else None
        items.append(
            {
                "id": item_id,
                "room": room,
                "roomSlug": slug(room),
                "category": label(categoria, CATEGORY_LABELS),
                "name": str(nombre).strip(),
                "description": str(descripcion).strip() if descripcion else None,
                "quantity": quantity,
                "measures": str(medidas).strip() if medidas else None,
                "condition": str(estado).strip() if estado else None,
                "price": price,
                "status": str(venta).strip() if venta else None,
                "image": image_path,
                **({"images": images} if len(images) > 1 else {}),
            }
        )

    payload = {
        "title": "Estancia La Taba",
        "subtitle": "Catálogo de muebles y objetos",
        "updated": "2026-09-24",
        "updatedLabel": "24 de septiembre de 2026",
        "rooms": [{"name": room, "slug": slug(room)} for room in room_order],
        "items": items,
    }
    (ROOT / "catalog.js").write_text(
        "window.CATALOG = " + json.dumps(payload, ensure_ascii=False, indent=2) + ";\n",
        encoding="utf-8",
    )
    export_hero()
    print(f"items {len(items)} photos {len(items) - missing} missing {missing} rooms {len(room_order)}")


if __name__ == "__main__":
    main()
